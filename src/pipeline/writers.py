from __future__ import annotations

import csv
from pathlib import Path

from src.pipeline.document import Document


def write_txt(doc: Document, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        for line in doc.lines:
            f.write(line.text + "\n")


def write_csv(doc: Document, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for line in doc.lines:
            writer.writerow([cell.text for cell in line.cells])


def write_xlsx(doc: Document, output_path: Path) -> None:
    import openpyxl

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    # Remove the default sheet
    wb.remove(wb.active)

    sheets: dict[str, openpyxl.worksheet.worksheet.Worksheet] = {}
    for line in doc.lines:
        sheet_name = line.cells[0].sheet_name if line.cells and line.cells[0].sheet_name else "Sheet"
        if sheet_name not in sheets:
            sheets[sheet_name] = wb.create_sheet(title=sheet_name)
        ws = sheets[sheet_name]
        ws.append([cell.text for cell in line.cells])

    wb.save(output_path)


def write_docx(doc: Document, output_path: Path) -> None:
    import docx

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for line in doc.lines:
        d.add_paragraph(line.text)
    d.save(output_path)


_WRITERS = {
    "txt": write_txt,
    "csv": write_csv,
    "xlsx": write_xlsx,
    "docx": write_docx,
}


def write_file(doc: Document, output_path: Path, preserve_format: bool) -> None:
    if preserve_format:
        fmt = doc.source_format
    else:
        fmt = "txt"
    writer = _WRITERS.get(fmt)
    if writer is None:
        raise ValueError(f"Unsupported output format: {fmt}")
    writer(doc, output_path)
